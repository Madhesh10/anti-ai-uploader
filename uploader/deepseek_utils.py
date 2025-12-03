# uploader/deepseek_utils.py
import os
import json
import logging
from io import BytesIO

import requests
from docx import Document as DocxReader
import pdfplumber
import openpyxl

logger = logging.getLogger(__name__)
logger.addHandler(logging.NullHandler())

# ---------------- DeepSeek config (read-from-env when used) ----------------
DEEPSEEK_BASE = os.getenv("DEEPSEEK_BASE", "https://api.deepseek.com")
DEFAULT_MODEL = os.getenv("DEEPSEEK_MODEL", "deepseek-chat")


def _get_key():
    """Return the DEEPSEEK_API_KEY from environment (or None)."""
    return os.getenv("DEEPSEEK_API_KEY")


def _mask_key(key: str | None) -> str:
    if not key:
        return "<NONE>"
    if len(key) <= 10:
        return key[:3] + "..." + key[-3:]
    return key[:6] + "..." + key[-4:]


def _require_key():
    if not _get_key():
        raise RuntimeError(
            "DEEPSEEK_API_KEY not set. Add it to Render environment variables or local .env."
        )


# ---------------- TEXT EXTRACTION ----------------
def extract_text_from_file(file_or_path):
    """
    Supports UploadedFile objects and file paths (str).
    Returns text (string).
    """
    # get bytes + name
    if isinstance(file_or_path, str):
        with open(file_or_path, "rb") as fh:
            b = fh.read()
        name = file_or_path.lower()
    else:
        try:
            file_or_path.seek(0)
        except Exception:
            pass
        b = file_or_path.read()
        name = getattr(file_or_path, "name", "").lower()

    bio = BytesIO(b)

    # DOCX
    if name.endswith(".docx"):
        bio.seek(0)
        doc = DocxReader(bio)
        return "\n".join(p.text for p in doc.paragraphs)

    # PDF
    if name.endswith(".pdf"):
        bio.seek(0)
        text = ""
        with pdfplumber.open(bio) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    text += t + "\n"
        return text

    # Excel
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        bio.seek(0)
        wb = openpyxl.load_workbook(bio, data_only=True)
        sheet = wb.active
        out = ""
        for row in sheet.iter_rows(values_only=True):
            out += " ".join(str(c) for c in row if c is not None) + "\n"
        return out

    # Fallback: treat as text
    try:
        return b.decode("utf-8", errors="ignore")
    except Exception:
        return str(b)


# ---------------- OPTIONAL: upload file to DeepSeek ----------------
def upload_file_to_deepseek(file_bytes, filename, metadata=None, timeout=30):
    """
    Upload file to DeepSeek. Returns dict:
      - success: JSON response from DeepSeek
      - error: {error info}
    """
    key = _get_key()
    if not key:
        return {"error": "no_api_key", "message": "DEEPSEEK_API_KEY not set"}

    url = f"{DEEPSEEK_BASE}/v1/files"
    headers = {"Authorization": f"Bearer {key}"}
    files = {"file": (filename, file_bytes)}
    data = {"purpose": "rag", "metadata": json.dumps(metadata or {})}

    try:
        r = requests.post(url, headers=headers, files=files, data=data, timeout=timeout)
    except requests.RequestException as e:
        logger.exception("upload_file_to_deepseek request failed")
        return {"error": "request_failed", "message": str(e), "endpoint": url}

    if r.status_code >= 400:
        body = _safe_json(r)
        return {
            "error": f"{r.status_code} {r.reason}",
            "body": body,
            "endpoint": url,
        }

    return _safe_json(r)


# ---------------- MAIN: ask DeepSeek with chat completions ----------------
def ask_deepseek(question, context=None, model=None, timeout=20):
    """
    Call DeepSeek chat model and return dict:
      - {"result": "<text>", "raw": <full_json>}
      - or {"error": "...", ...}
    """
    key = _get_key()
    if not key:
        return {"error": "no_api_key", "message": "DEEPSEEK_API_KEY not set"}

    model = model or DEFAULT_MODEL
    url = f"{DEEPSEEK_BASE}/v1/chat/completions"
    headers = {
        "Authorization": f"Bearer {key}",
        "Content-Type": "application/json",
    }

    if context:
        user_content = f"Context:\n{context}\n\nQuestion: {question}"
    else:
        user_content = question

    payload = {
        "model": model,
        "messages": [
            {
                "role": "system",
                "content": "You are an AI assistant that answers strictly using the provided context.",
            },
            {"role": "user", "content": user_content},
        ],
    }

    try:
        r = requests.post(url, headers=headers, json=payload, timeout=timeout)
    except requests.RequestException as e:
        logger.exception("ask_deepseek request failed")
        return {"error": "request_failed", "message": str(e), "endpoint": url}

    if r.status_code >= 400:
        body = _safe_json(r)
        return {"error": f"{r.status_code} {r.reason}", "body": body, "endpoint": url}

    data = _safe_json(r)
    if isinstance(data, dict):
        assistant_text = _extract_assistant_text(data)
        if assistant_text:
            return {"result": assistant_text, "raw": data}
        else:
            return {"error": "no_text_in_response", "raw": data}
    else:
        return {"error": "invalid_json", "body": data}


def _safe_json(response):
    """Return response.json() safely, fallback to text."""
    try:
        return response.json()
    except Exception:
        return {"text": response.text}


def _extract_assistant_text(data):
    """Try standard deepseek / openai-like response shapes."""
    try:
        choices = data.get("choices")
        if not choices:
            return None
        first = choices[0]
        # modern shape: choices[0]["message"]["content"]
        msg = first.get("message")
        if msg and "content" in msg:
            return msg["content"]
        # fallback shape: choices[0].get("text")
        if "text" in first:
            return first.get("text")
    except Exception:
        return None
    return None


# ---------------- Helper: verify key (for quick local tests) ----------------
def verify_key(timeout=10):
    """
    Quick verification helper: attempts a minimal request to check the key.
    Returns {'ok': True} on success or error dict on failure.
    """
    key = _get_key()
    if not key:
        return {"ok": False, "error": "no_api_key"}

    url = f"{DEEPSEEK_BASE}/v1/chat/completions"
    headers = {"Authorization": f"Bearer {key}", "Content-Type": "application/json"}
    payload = {"model": DEFAULT_MODEL, "messages": [{"role": "user", "content": "ping"}]}

    try:
        r = requests.post(url, headers=headers, json=payload, timeout=timeout)
    except requests.RequestException as e:
        return {"ok": False, "error": "request_failed", "message": str(e)}

    if r.status_code == 200:
        return {"ok": True, "status_code": 200}
    else:
        return {"ok": False, "status_code": r.status_code, "body": _safe_json(r)}
